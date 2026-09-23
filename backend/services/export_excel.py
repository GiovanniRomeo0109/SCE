"""
services/export_excel.py — Esportazioni Excel del blocco 4

1. Cronoprogramma: Gantt settimanale con date reali (dalla data di inizio lavori),
   uomini-giorno per fase e foglio di confronto dei due metodi.
2. Costi della sicurezza: stima analitica per categoria dell'Allegato XV 4.1.1 con subtotali e totale.

Tutti i calcoli sono FORMULE (si aggiornano se il CSP modifica un valore in Excel).
Celle da compilare/modificabili: testo blu su fondo giallo chiaro. Font Arial.
"""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

FONT = "Arial"
BLU = Font(name=FONT, size=10, color="0000FF")
NERO = Font(name=FONT, size=10)
GRASSETTO = Font(name=FONT, size=10, bold=True)
TITOLO = Font(name=FONT, size=14, bold=True, color="1A3A5C")
BIANCO = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_INTEST = PatternFill("solid", fgColor="1A3A5C")
F_INPUT = PatternFill("solid", fgColor="FFF8E1")
F_CAT = PatternFill("solid", fgColor="EEF4FA")
F_BARRA = PatternFill("solid", fgColor="C88B2A")
F_TOT = PatternFill("solid", fgColor="E8F5E9")
SOTTILE = Side(style="thin", color="C8D2DC")
BORDO = Border(left=SOTTILE, right=SOTTILE, top=SOTTILE, bottom=SOTTILE)
EURO = '€ #,##0.00;-€ #,##0.00;"-"'
DATA = "dd/mm/yyyy"
A_CAPO = Alignment(wrap_text=True, vertical="top")
CENTRO = Alignment(horizontal="center", vertical="center")


def _cella(ws, rif, valore, font=NERO, fill=None, fmt=None, align=None, bordo=True):
    c = ws[rif]
    c.value = valore
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = align
    if bordo:
        c.border = BORDO
    return c


def _intestazione(ws, riga, titoli, larghezze=None):
    for i, t in enumerate(titoli, start=1):
        _cella(ws, f"{L(i)}{riga}", t, BIANCO, F_INTEST, align=Alignment(wrap_text=True, vertical="center", horizontal="center"))
        if larghezze and i <= len(larghezze) and larghezze[i - 1]:
            ws.column_dimensions[L(i)].width = larghezze[i - 1]


def _salva(wb) -> bytes:
    wb.calculation.fullCalcOnLoad = True      # Excel ricalcola all'apertura
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ══════════════════════════════════════════════════════════════════════════════
# 1. Cronoprogramma
# ══════════════════════════════════════════════════════════════════════════════

def cronoprogramma_xlsx(nome_progetto: str, ug: dict) -> bytes:
    righe = ug["cronoprogramma"]["righe"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Cronoprogramma"

    _cella(ws, "A1", f"Cronoprogramma dei lavori — {nome_progetto}", TITOLO, bordo=False)
    _cella(ws, "A3", "Data di inizio lavori", GRASSETTO, bordo=False)
    inizio = ug.get("data_inizio_lavori")
    try:
        d0 = date.fromisoformat(inizio) if inizio else None
    except ValueError:
        d0 = None
    _cella(ws, "C3", d0, BLU, F_INPUT, DATA)
    _cella(ws, "A4", "Giorni lavorativi per settimana", GRASSETTO, bordo=False)
    _cella(ws, "C4", ug["giorni_per_settimana"], BLU, F_INPUT, "0")
    _cella(ws, "E3", "Le celle gialle con testo blu sono modificabili: date, durate e uomini-giorno si ricalcolano.",
           Font(name=FONT, size=9, italic=True, color="5A6B7D"), bordo=False)
    if not d0:
        _cella(ws, "E4", "Data di inizio non indicata nel progetto: inseriscila in C3 per calcolare le date.",
               Font(name=FONT, size=9, italic=True, color="C0392B"), bordo=False)

    fisse = ["Cod.", "Fase / lavorazione", "Impresa", "Addetti medi", "Sett. inizio", "Durata (sett.)",
             "Sett. fine", "Dal", "Al", "Uomini-giorno", "Contemporanea a"]
    R0 = 6
    _intestazione(ws, R0, fisse, [7, 38, 24, 9, 8, 9, 8, 11, 11, 11, 18])
    ws.row_dimensions[R0].height = 32

    max_sett = max([int(r["fine"]) for r in righe if r["fine"]] + [int((r["inizio"] or 1) + (r["durata"] or 1) - 1) for r in righe] + [4])
    max_sett = min(max_sett + 2, 156)
    c0 = len(fisse) + 1
    for k in range(max_sett):
        col = L(c0 + k)
        _cella(ws, f"{col}{R0}", k + 1, BIANCO, F_INTEST, "0", CENTRO)
        # riga sopra l'intestazione: lunedì della settimana
        _cella(ws, f"{col}{R0 - 1}", f'=IF($C$3="","",$C$3-WEEKDAY($C$3,2)+1+({col}{R0}-1)*7)',
               Font(name=FONT, size=7, color="5A6B7D"), fmt="dd/mm", align=Alignment(text_rotation=90, horizontal="center"))
        ws.column_dimensions[col].width = 3.2
    ws.row_dimensions[R0 - 1].height = 38

    r = R0 + 1
    for riga in righe:
        _cella(ws, f"A{r}", riga["cod"], NERO)
        _cella(ws, f"B{r}", riga["fase"], NERO, align=A_CAPO)
        _cella(ws, f"C{r}", riga["impresa"], NERO, align=A_CAPO)
        _cella(ws, f"D{r}", riga["addetti"], BLU, F_INPUT, "0")
        _cella(ws, f"E{r}", riga["inizio"], BLU, F_INPUT, "0")
        _cella(ws, f"F{r}", riga["durata"], BLU, F_INPUT, "0.0")
        _cella(ws, f"G{r}", f'=IF(AND(ISNUMBER(E{r}),ISNUMBER(F{r})),E{r}+ROUNDUP(F{r},0)-1,"")', NERO, fmt="0")
        _cella(ws, f"H{r}", f'=IF(OR($C$3="",NOT(ISNUMBER(E{r}))),"",$C$3-WEEKDAY($C$3,2)+1+(E{r}-1)*7)', NERO, fmt=DATA)
        _cella(ws, f"I{r}", f'=IF(OR(H{r}="",NOT(ISNUMBER(G{r}))),"",$C$3-WEEKDAY($C$3,2)+1+(G{r}-1)*7+$C$4-1)', NERO, fmt=DATA)
        _cella(ws, f"J{r}", f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(F{r})),F{r}*$C$4*D{r},"")', NERO, fmt="#,##0.0")
        _cella(ws, f"K{r}", riga["contemporanea"], NERO, align=A_CAPO)
        for k in range(max_sett):
            ws[f"{L(c0 + k)}{r}"].border = BORDO
        r += 1
    ultima = r - 1

    if righe:
        area = f"{L(c0)}{R0 + 1}:{L(c0 + max_sett - 1)}{ultima}"
        prima_col = L(c0)
        ws.conditional_formatting.add(area, FormulaRule(
            formula=[f"AND(ISNUMBER($E{R0 + 1}),ISNUMBER($G{R0 + 1}),{prima_col}${R0}>=$E{R0 + 1},{prima_col}${R0}<=$G{R0 + 1})"],
            fill=F_BARRA))
    _cella(ws, f"I{r + 1}", "Totale uomini-giorno", GRASSETTO, F_TOT)
    _cella(ws, f"J{r + 1}", f"=SUM(J{R0 + 1}:J{max(ultima, R0 + 1)})", GRASSETTO, F_TOT, "#,##0.0")
    ws.freeze_panes = ws[f"{L(c0)}{R0 + 1}"]
    tot_crono = f"Cronoprogramma!$J${r + 1}"

    # ── Foglio uomini-giorno ──
    u = wb.create_sheet("Uomini-giorno")
    inc = ug["incidenza"]
    u.column_dimensions["A"].width = 52
    u.column_dimensions["B"].width = 18
    u.column_dimensions["C"].width = 60
    _cella(u, "A1", "Uomini-giorno — confronto dei due metodi", TITOLO, bordo=False)
    _cella(u, "A3", "Metodo A — dal cronoprogramma", GRASSETTO, F_CAT)
    _cella(u, "A4", "Σ durata (settimane) × giorni/settimana × addetti medi")
    _cella(u, "B4", f"={tot_crono}", Font(name=FONT, size=10, color="008000"), fmt="#,##0.0")
    _cella(u, "A6", "Metodo B — dall'incidenza della manodopera", GRASSETTO, F_CAT)
    _cella(u, "A7", "Importo dei lavori (€)")
    _cella(u, "B7", inc["importo_lavori"], BLU, F_INPUT, EURO)
    _cella(u, "A8", "Incidenza della manodopera (%)")
    _cella(u, "B8", (inc["incidenza_csp"] / 100) if inc["incidenza_csp"] is not None else
           ((inc["incidenza_documenti"] / 100) if inc["incidenza_documenti"] else None), BLU, F_INPUT, "0.0%")
    _cella(u, "A9", "Importo manodopera da computo (€) — se presente prevale")
    man_computo = inc["importo_manodopera"] if (inc["origine_manodopera"] or "").startswith("computo (importo") else None
    _cella(u, "B9", man_computo, BLU, F_INPUT, EURO)
    _cella(u, "A10", "Importo della manodopera (€)")
    _cella(u, "B10", '=IF(ISNUMBER(B9),B9,IF(AND(ISNUMBER(B7),ISNUMBER(B8)),B7*B8,""))', NERO, fmt=EURO)
    _cella(u, "A11", "Costo giornaliero di un lavoratore (€/giorno)")
    _cella(u, "B11", inc["costo_giornaliero"], BLU, F_INPUT, EURO)
    _cella(u, "A12", "Uomini-giorno (importo manodopera ÷ costo giornaliero)")
    _cella(u, "B12", '=IF(AND(ISNUMBER(B10),ISNUMBER(B11),B11>0),B10/B11,"")', NERO, fmt="#,##0.0")
    prop = inc.get("proposta")
    if prop:
        _cella(u, "C11", f"Proposto: media di {len(prop['qualifiche'])} qualifiche edili ({prop['media_oraria']:.2f} €/h) × 8 h — "
                         f"elenco \"{prop['elenco']}\"", Font(name=FONT, size=9, italic=True, color="5A6B7D"), bordo=False)
    if inc["incidenza_csp"] is None and not man_computo and not inc["incidenza_documenti"]:
        _cella(u, "C8", "Inserisci la percentuale di incidenza della manodopera per calcolare questo metodo",
               Font(name=FONT, size=9, italic=True, color="C0392B"), bordo=False)

    _cella(u, "A14", "Valore per il PSC", GRASSETTO, F_TOT)
    scelta = ug.get("scelta")
    _cella(u, "A15", "Metodo scelto dal CSP (A = cronoprogramma, B = incidenza)")
    _cella(u, "B15", "A" if scelta == "cronoprogramma" else "B" if scelta == "incidenza" else "", BLU, F_INPUT, align=CENTRO)
    _cella(u, "A16", "Uomini-giorno presunti")
    _cella(u, "B16", '=IF(B15="A",B4,IF(B15="B",B12,MAX(N(B4),N(B12))))', GRASSETTO, F_TOT, "#,##0.0")
    _cella(u, "C16", "Se il metodo non è indicato vale il maggiore dei due. Soglia di 200 uomini-giorno: "
                     "art. 99, co. 1, lett. c), D.Lgs. 81/2008.", Font(name=FONT, size=9, italic=True, color="5A6B7D"), bordo=False)
    return _salva(wb)


# ══════════════════════════════════════════════════════════════════════════════
# 2. Costi della sicurezza
# ══════════════════════════════════════════════════════════════════════════════

def costi_xlsx(nome_progetto: str, righe: list, categorie: dict, escluse: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Costi sicurezza"
    _cella(ws, "A1", f"Stima analitica dei costi della sicurezza — {nome_progetto}", TITOLO, bordo=False)
    _cella(ws, "A2", "D.Lgs. 81/2008, Allegato XV, punto 4.1.1 — costi non soggetti a ribasso. Importi IVA esclusa.",
           Font(name=FONT, size=9, italic=True, color="5A6B7D"), bordo=False)
    _cella(ws, "A3", "Quantità e prezzi in giallo sono modificabili; gli importi si ricalcolano. "
                     "Le righe senza prezzo sono \"prezzo da definire\".",
           Font(name=FONT, size=9, italic=True, color="5A6B7D"), bordo=False)
    titoli = ["Cat.", "Misura", "Codice voce", "Descrizione voce", "Elenco", "U.M. voce", "Quantità",
              "Prezzo unitario", "Importo", "Note"]
    R0 = 5
    _intestazione(ws, R0, titoli, [6, 40, 17, 52, 20, 9, 10, 13, 14, 36])
    ws.row_dimensions[R0].height = 28

    r = R0 + 1
    subtotali = []
    gruppi = {}
    for x in righe:
        gruppi.setdefault(x["categoria_lettera"], []).append(x)
    for cat in categorie["categorie"]:
        k = cat["lettera"]
        voci = gruppi.get(k, [])
        if not voci:
            continue
        etichetta = f"{k}) {cat['descrizione']}" if k != "altro" else "Altre misure (categoria non indicata)"
        _cella(ws, f"A{r}", etichetta, GRASSETTO, F_CAT)
        ws.merge_cells(f"A{r}:J{r}")
        r += 1
        prima = r
        for x in voci:
            _cella(ws, f"A{r}", k, NERO, align=CENTRO)
            _cella(ws, f"B{r}", x["misura"], NERO, align=A_CAPO)
            _cella(ws, f"C{r}", x["codice"] or "", NERO)
            _cella(ws, f"D{r}", x["descrizione_voce"] or "Prezzo da definire", NERO if x["descrizione_voce"] else
                   Font(name=FONT, size=10, italic=True, color="C0392B"), align=A_CAPO)
            livello = {"progetto": "Progetto", "account": "Account", "sistema": "Sistema", "manuale": "Manuale"}.get(x["livello"] or "", "")
            _cella(ws, f"E{r}", f"{livello}: {x['elenco_nome']}" if x.get("elenco_nome") else livello, NERO, align=A_CAPO)
            _cella(ws, f"F{r}", x["um_voce"] or x["um_misura"] or "", NERO, align=CENTRO)
            _cella(ws, f"G{r}", x["quantita"], BLU, F_INPUT, "#,##0.00")
            _cella(ws, f"H{r}", x["prezzo"], BLU, F_INPUT, EURO)
            _cella(ws, f"I{r}", f'=IF(AND(ISNUMBER(G{r}),ISNUMBER(H{r})),G{r}*H{r},0)', NERO, fmt=EURO)
            nota = x.get("nota") or ""
            if x.get("um_diverse"):
                nota = (f"U.M. diverse: misura in {x['um_misura']}, voce in {x['um_voce']}. " + nota).strip()
            _cella(ws, f"J{r}", nota, Font(name=FONT, size=9, color="5A6B7D"), align=A_CAPO)
            r += 1
        _cella(ws, f"H{r}", f"Subtotale {k})" if k != "altro" else "Subtotale", GRASSETTO)
        _cella(ws, f"I{r}", f"=SUM(I{prima}:I{r - 1})", GRASSETTO, fmt=EURO)
        subtotali.append(f"I{r}")
        r += 2
    _cella(ws, f"H{r}", "TOTALE COSTI DELLA SICUREZZA", GRASSETTO, F_TOT, align=A_CAPO)
    _cella(ws, f"I{r}", "=" + "+".join(subtotali) if subtotali else 0, GRASSETTO, F_TOT, EURO)
    ws.row_dimensions[r].height = 28
    da_def = sum(1 for x in righe if x["prezzo"] is None or x["quantita"] is None)
    if da_def:
        _cella(ws, f"B{r}", f"Attenzione: {da_def} misure con prezzo o quantità da definire non sono conteggiate.",
               Font(name=FONT, size=10, bold=True, color="C0392B"), bordo=False)
    ws.freeze_panes = ws[f"A{R0 + 1}"]

    e = wb.create_sheet("Voci escluse")
    _cella(e, "A1", "Voci escluse: costi ordinari delle imprese (non rientrano nei costi della sicurezza del PSC)", TITOLO, bordo=False)
    _intestazione(e, 3, ["Voce", "Motivo dell'esclusione"], [60, 80])
    for i, (voce, motivo) in enumerate(escluse, start=4):
        _cella(e, f"A{i}", voce, NERO, align=A_CAPO)
        _cella(e, f"B{i}", motivo, NERO, align=A_CAPO)
    return _salva(wb)

"""
services/psc_docx.py — Documento finale del PSC (blocco 6)

Struttura: copertina (logo e dati dello studio, cantiere, committente, CSP, revisione) →
riepilogo dei punti DA VERIFICARE → indice → 12 capitoli (punto 23 della metodologia) → firma.
Intestazione con logo e studio; piè di pagina "Rev. N del gg/mm/aaaa — Pag. X di Y".
Il controllo di coerenza (tappa 12) e la checklist A–M restano nell'app, non nel documento.
"""
import json
import os
import re
import uuid
from datetime import date

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from database import get_conn, cartella_documenti

BLU = RGBColor(0x1A, 0x3A, 0x5C)
GRIGIO = RGBColor(0x5A, 0x6B, 0x7D)
DAV = "DA VERIFICARE"
LARGHEZZA_UTILE_CM = 17.0

CAPITOLI = [
    (1, "Identificazione del cantiere"), (2, "Descrizione dell'opera"), (3, "Descrizione del contesto"),
    (4, "Organizzazione del cantiere"), (5, "Lavorazioni"), (6, "Analisi dei rischi"),
    (7, "Misure preventive e protettive"), (8, "Coordinamento"), (9, "Cronoprogramma"),
    (10, "Emergenze"), (11, "Costi della sicurezza"), (12, "Planimetrie e allegati"),
]


# ══════════════════════════════════════════════════════════════════════════════
# Raccolta dei dati
# ══════════════════════════════════════════════════════════════════════════════

def _tappe(progetto_id):
    conn = get_conn()
    try:
        rows = conn.execute("SELECT numero, contenuto_json FROM progetto_tappe WHERE progetto_id=?", (progetto_id,)).fetchall()
    finally:
        conn.close()
    return {r["numero"]: json.loads(r["contenuto_json"]) for r in rows if r["contenuto_json"]}


def _sez(tappe, n, sid):
    return next((s for s in tappe.get(n, {}).get("sezioni", []) if s["id"] == sid), None)


def _euro(v):
    return "—" if v is None else f"€ {v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _num(v, dec=2):
    if v is None:
        return "—"
    s = f"{v:,.{dec}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return s.rstrip("0").rstrip(",") if "," in s else s


def costruisci_piano(progetto_id: int, username: str, coordinatore: dict, studio: dict) -> dict:
    """Blocchi di ogni capitolo: ('testo', titolo, testo) · ('tabella', titolo, colonne, righe) ·
    ('gantt', righe, n_settimane) · ('costi', righe, totali) · ('immagine', titolo, png) · ('nota', testo)."""
    from services import costi_sicurezza as cs, uomini_giorno as ugm, presidi as pres, schemi_render
    t = _tappe(progetto_id)
    conn = get_conn()
    try:
        p = dict(conn.execute("SELECT * FROM progetti WHERE id=?", (progetto_id,)).fetchone())
        schemi = [dict(r) for r in conn.execute(
            "SELECT * FROM progetto_schemi WHERE progetto_id=? AND includi_psc=1 ORDER BY id", (progetto_id,)).fetchall()]
        documenti = [dict(r) for r in conn.execute(
            """SELECT nome_file, COALESCE(tipo_csp, tipo_ai) AS tipo, pagine FROM progetto_documenti
               WHERE progetto_id=? AND stato='completato' ORDER BY priorita, id""", (progetto_id,)).fetchall()]
    finally:
        conn.close()

    def testo(n, sid):
        s = _sez(t, n, sid)
        return [("testo", s["titolo"], s["testo"])] if s and s.get("testo") else []

    def tabella(n, sid, titolo=None):
        s = _sez(t, n, sid)
        return [("tabella", titolo or s["titolo"], s["colonne"], s["righe"])] if s and s.get("righe") else []

    ug = ugm.calcola(progetto_id, username)
    crono = ug["cronoprogramma"]["righe"]
    durata = max([r["fine"] for r in crono if r.get("fine")] or [0])
    piano = {}

    # 1. Identificazione del cantiere
    csp = []
    if coordinatore:
        nome = f"{coordinatore.get('nome', '')} {coordinatore.get('cognome', '')}".strip()
        iscr = " ".join(x for x in (coordinatore.get("ordine_professionale"), coordinatore.get("provincia_ordine"),
                                    f"n. {coordinatore['numero_ordine']}" if coordinatore.get("numero_ordine") else "") if x)
        csp = [["Coordinatore per la progettazione (CSP)", nome,
                "; ".join(x for x in (iscr, coordinatore.get("pec") or coordinatore.get("email"), coordinatore.get("telefono")) if x)]]
    s_sog = _sez(t, 3, "soggetti")
    righe_sog = [r for r in (s_sog["righe"] if s_sog else []) if not (csp and "csp" in r[0].lower())]
    piano[1] = ([("tabella", "Ubicazione del cantiere", ["Dato", "Valore"],
                  [["Cantiere", p["nome"]], ["Indirizzo", p.get("indirizzo_cantiere") or DAV]])]
                + ([("tabella", "Soggetti con compiti di sicurezza", ["Ruolo", "Nominativo", "Recapiti / riferimenti"],
                     csp + righe_sog)] if (csp or righe_sog) else [])
                + tabella(3, "imprese") + testo(3, "note_soggetti"))

    # 2. Descrizione dell'opera
    importo = ug["incidenza"]["importo_lavori"]
    dati_ec = [["Importo dei lavori", _euro(importo) if importo else DAV],
               ["Costi della sicurezza (stima analitica, cap. 11)", _euro(cs.totali(cs.righe(progetto_id))["totale"])],
               ["Durata presunta", f"{_num(durata, 0)} settimane" if durata else DAV],
               ["Data di inizio lavori", date.fromisoformat(p["data_inizio_lavori"]).strftime("%d/%m/%Y") if p.get("data_inizio_lavori")
                else "Non indicata: il cronoprogramma è espresso in giorni di cantiere (1° giorno = lunedì, settimana di 5 giorni lavorativi)"],
               ["Entità presunta del cantiere", f"{_num(ug['valore'], 1)} uomini-giorno" if ug.get("valore") else DAV]]
    piano[2] = (testo(1, "descrizione") + tabella(1, "dati_opera")
                + [("tabella", "Dati economici e temporali", ["Dato", "Valore"], dati_ec)]
                + testo(1, "analisi_progetto") + tabella(1, "scelte_progettuali"))
    # 3-5
    piano[3] = testo(2, "inquadramento") + tabella(2, "caratteristiche_area") + tabella(2, "rischi_esterni")
    piano[4] = testo(9, "layout") + tabella(9, "elementi_cantiere")
    if schemi:
        piano[4].append(("nota", f"Lo schema di organizzazione del cantiere è riportato nel capitolo 12 ({len(schemi)} tavole)."))
    piano[5] = tabella(4, "lavorazioni") + testo(4, "note_lavorazioni")
    # 6. Rischi (area, lavorazioni, particolari, interferenziali)
    piano[6] = (tabella(6, "rischi_area") + tabella(6, "rischi_lavorazioni") + tabella(6, "rischi_particolari")
                + testo(6, "amianto") + tabella(7, "interferenze", "Rischi interferenziali") + testo(7, "note_interferenze"))
    # 7-8
    piano[7] = testo(8, "principi") + tabella(8, "misure_interferenze")
    piano[8] = tabella(10, "uso_comune") + tabella(10, "attivita_coordinamento") + testo(10, "cooperazione")
    # 9. Cronoprogramma
    piano[9] = testo(5, "impostazione") + tabella(5, "cronoprogramma")
    if crono and durata:
        piano[9].append(("gantt", crono, int(durata)))
    metodo = {"cronoprogramma": "dal cronoprogramma (settimane × 5 giorni × addetti medi)",
              "incidenza": "dall'incidenza della manodopera"}.get(ug.get("scelta"), "")
    piano[9].append(("tabella", "Uomini-giorno", ["Metodo", "Uomini-giorno"],
                     [["A — dal cronoprogramma", _num(ug["cronoprogramma"]["uomini_giorno"], 1)],
                      ["B — dall'incidenza della manodopera", _num(ug["incidenza"]["uomini_giorno"], 1)],
                      [f"Valore adottato nel PSC ({metodo})" if metodo else "Valore adottato nel PSC",
                       _num(ug.get("valore"), 1) if ug.get("valore") else DAV]]))
    # 10. Emergenze
    righe_pres = pres.righe_tabella(progetto_id)
    piano[10] = testo(9, "emergenze") + ([("tabella", "Presidi e numeri di emergenza",
                                           ["Presidio", "Riferimento / numero", "Distanza o percorso indicativo"], righe_pres)]
                                         if righe_pres else tabella(9, "presidi"))
    # 11. Costi
    righe_costi = cs.righe(progetto_id)
    piano[11] = ([("costi", righe_costi, cs.totali(righe_costi))] if righe_costi else
                 [("nota", "Stima analitica dei costi della sicurezza non ancora compilata: " + DAV)])
    piano[11] += tabella(11, "costi_ordinari") + testo(11, "criteri_stima")
    # 12. Planimetrie e allegati
    piano[12] = [("immagine", s["nome"], schemi_render.disegna(s)) for s in schemi]
    if not schemi:
        piano[12].append(("nota", "Nessuno schema di cantiere selezionato per il PSC: " + DAV))
    if documenti:
        from services.documenti_progetto import etichetta
        piano[12].append(("tabella", "Documentazione di progetto esaminata", ["Documento", "Tipo", "Pagine"],
                          [[d["nome_file"], etichetta(d["tipo"]), str(d["pagine"] or "")] for d in documenti
                           if d["tipo"] != "elenco_prezzi"]))
    return {"progetto": p, "piano": piano, "tappe": t}


def punti_da_verificare(piano: dict) -> list:
    out = []
    for n, titolo in CAPITOLI:
        for b in piano.get(n, []):
            if b[0] == "testo" and DAV in b[2].upper():
                out.append((n, titolo, b[1]))
            elif b[0] == "tabella":
                k = sum(1 for r in b[3] for c in r if DAV in (c or "").upper())
                if k:
                    out.append((n, titolo, f"{b[1]} ({k} {'dato' if k == 1 else 'dati'})"))
            elif b[0] == "nota" and DAV in b[1]:
                out.append((n, titolo, b[1].split(":")[0]))
            elif b[0] == "costi":
                k = b[2]["da_definire"]
                if k:
                    out.append((n, titolo, f"{k} {'voce' if k == 1 else 'voci'} con prezzo da definire"))
    return out


# ══════════════════════════════════════════════════════════════════════════════
# Scrittura DOCX
# ══════════════════════════════════════════════════════════════════════════════

def _run_dav(par, testo, grassetto=False, dim=None, colore=None):
    """Scrive il testo evidenziando in giallo ogni 'DA VERIFICARE'."""
    for parte in re.split(r"(DA VERIFICARE)", testo or "", flags=re.IGNORECASE):
        if not parte:
            continue
        r = par.add_run(parte)
        r.bold = grassetto or None
        if dim:
            r.font.size = Pt(dim)
        if colore:
            r.font.color.rgb = colore
        if parte.upper() == DAV:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            r.bold = True


def _ombreggia(cella, esadecimale):
    tcPr = cella._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), esadecimale)
    tcPr.append(shd)


def _campo(par, istruzione):
    """Campo Word (PAGE, NUMPAGES) aggiornato automaticamente."""
    r = par.add_run()
    for tipo, testo in (("begin", None), (None, istruzione), ("separate", None), (None, "1"), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tipo)
        elif testo == istruzione:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = f" {istruzione} "
        else:
            el = OxmlElement("w:t"); el.text = testo
        r._r.append(el)


def _larghezze(colonne, righe, totale=LARGHEZZA_UTILE_CM):
    pesi = []
    for i, c in enumerate(colonne):
        lung = max([len(c)] + [len(r[i]) for r in righe[:40] if i < len(r)])
        pesi.append(min(max(lung, 6), 60))
    s = sum(pesi)
    return [totale * p / s for p in pesi]


def _tabella(doc, colonne, righe, dim=8.5, larghezze=None, intestazione=True):
    tab = doc.add_table(rows=1, cols=len(colonne))
    tab.style = "Table Grid"
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    larghezze = larghezze or _larghezze(colonne, righe)
    # Larghezze anche sulla griglia della tabella (w:tblGrid): LibreOffice e Word la usano per impaginare
    for i, col in enumerate(tab.columns):
        col.width = Cm(larghezze[i])
    tblPr = tab._tbl.tblPr
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); tblPr.append(layout)
    for i, c in enumerate(colonne):
        cella = tab.rows[0].cells[i]
        cella.width = Cm(larghezze[i])
        cella.paragraphs[0].text = ""
        _run_dav(cella.paragraphs[0], c, grassetto=True, dim=dim, colore=RGBColor(0xFF, 0xFF, 0xFF))
        _ombreggia(cella, "1A3A5C")
    # l'intestazione si ripete su ogni pagina
    trPr = tab.rows[0]._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)
    for r in righe:
        celle = tab.add_row().cells
        for i in range(len(colonne)):
            celle[i].width = Cm(larghezze[i])
            _run_dav(celle[i].paragraphs[0], r[i] if i < len(r) else "", dim=dim)
    if not intestazione:
        tab._tbl.remove(tab.rows[0]._tr)
    return tab


def _sottotitolo(doc, testo):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    _run_dav(p, testo, grassetto=True, dim=11, colore=BLU)
    p.paragraph_format.keep_with_next = True


def _gantt(doc, righe, n_settimane):
    passo = 1 if n_settimane <= 26 else 2 if n_settimane <= 52 else 4
    colonne = list(range(1, n_settimane + 1, passo))
    fisse = 5.4
    w_sett = (LARGHEZZA_UTILE_CM - fisse) / max(len(colonne), 1)
    intest = ["Cod.", "Fase"] + [str(s) for s in colonne]
    larghezze = [1.1, fisse - 1.1] + [w_sett] * len(colonne)
    corpo = [[r["cod"] or "", r["fase"] or ""] + [""] * len(colonne) for r in righe]
    tab = _tabella(doc, intest, corpo, dim=6.5, larghezze=larghezze)
    for i, r in enumerate(righe, start=1):
        if not (r.get("inizio") and r.get("fine")):
            continue
        for j, s in enumerate(colonne):
            if r["inizio"] <= s + passo - 1 and s <= r["fine"]:
                _ombreggia(tab.rows[i].cells[2 + j], "C88B2A")
    if passo > 1:
        doc.add_paragraph().add_run(f"Ogni colonna rappresenta {passo} settimane.").font.size = Pt(8)


def _costi(doc, righe, totali):
    intest = ["Cat.", "Misura", "Voce di prezzo", "U.M.", "Quantità", "Prezzo unit.", "Importo"]
    corpo, marcature = [], []
    gruppi = {}
    for r in righe:
        gruppi.setdefault(r["categoria_lettera"], []).append(r)
    for cat in totali["categorie"]:
        k = cat["lettera"]
        corpo.append([f"{k})" if k != "altro" else "", cat["descrizione"], "", "", "", "", ""]); marcature.append("cat")
        for r in gruppi.get(k, []):
            descr = r.get("descrizione_voce") or ""
            descr = descr if len(descr) <= 180 else descr[:177].rsplit(" ", 1)[0] + "…"
            voce = f"{r['codice']} — {descr}" if r.get("codice") else (descr or "Prezzo da definire — " + DAV)
            corpo.append([k, r["misura"], voce, r.get("um_voce") or r.get("um_misura") or "",
                          _num(r["quantita"]), _euro(r["prezzo"]) if r["prezzo"] is not None else DAV,
                          _euro(r["importo"]) if r["importo"] is not None else DAV]); marcature.append("voce")
        corpo.append(["", f"Subtotale {k})" if k != "altro" else "Subtotale", "", "", "", "", _euro(cat["importo"])]); marcature.append("sub")
    corpo.append(["", "TOTALE COSTI DELLA SICUREZZA (non soggetti a ribasso)", "", "", "", "", _euro(totali["totale"])]); marcature.append("tot")
    tab = _tabella(doc, intest, corpo, dim=8, larghezze=[1.0, 4.3, 5.4, 1.2, 1.5, 1.8, 1.8])
    for i, m in enumerate(marcature, start=1):
        if m in ("cat", "sub", "tot"):
            for c in tab.rows[i].cells:
                _ombreggia(c, {"cat": "EEF4FA", "sub": "F4F6F9", "tot": "E8F5E9"}[m])
                for run in c.paragraphs[0].runs:
                    run.bold = True


def scrivi(dati: dict, studio: dict, coordinatore: dict, firma_path: str, logo_path: str,
           revisione: int, data_rev: str) -> bytes:
    p, piano = dati["progetto"], dati["piano"]
    rev = f"Rev. {revisione} del {data_rev}"
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for nome, dim in (("Heading 1", 14), ("Heading 2", 11)):
        h = doc.styles[nome]
        h.font.name = "Arial"; h.font.size = Pt(dim); h.font.bold = True; h.font.color.rgb = BLU

    sez = doc.sections[0]
    sez.page_height, sez.page_width = Cm(29.7), Cm(21.0)
    for lato in ("left_margin", "right_margin"):
        setattr(sez, lato, Cm(2.0))
    sez.top_margin, sez.bottom_margin = Cm(2.2), Cm(2.0)
    sez.different_first_page_header_footer = True

    # Intestazione e piè di pagina (non in copertina)
    hp = sez.header.paragraphs[0]
    if logo_path and os.path.exists(logo_path):
        hp.add_run().add_picture(logo_path, height=Cm(0.9))
        hp.add_run("   ")
    r = hp.add_run(f"{studio.get('nome') or ''}  —  PSC: {p['nome']}".strip(" —"))
    r.font.size = Pt(8); r.font.color.rgb = GRIGIO
    fp = sez.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(f"{rev}  —  Pag. "); r.font.size = Pt(8)
    _campo(fp, "PAGE")
    r = fp.add_run(" di "); r.font.size = Pt(8)
    _campo(fp, "NUMPAGES")

    # ── Copertina ──
    if logo_path and os.path.exists(logo_path):
        c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.add_run().add_picture(logo_path, height=Cm(2.5))
    if studio.get("nome"):
        c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(studio["nome"]); r.bold = True; r.font.size = Pt(12)
        dati_st = " · ".join(x for x in (studio.get("indirizzo"), f"P.IVA {studio['piva']}" if studio.get("piva") else "",
                                         studio.get("telefono"), studio.get("email")) if x)
        if dati_st:
            c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = c.add_run(dati_st); r.font.size = Pt(9); r.font.color.rgb = GRIGIO
    for _ in range(4):
        doc.add_paragraph()
    for testo, dim in (("PIANO DI SICUREZZA E COORDINAMENTO", 22), ("art. 100 e Allegato XV — D.Lgs. 9 aprile 2008, n. 81", 11)):
        c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(testo); r.bold = dim > 12; r.font.size = Pt(dim); r.font.color.rgb = BLU
    for _ in range(3):
        doc.add_paragraph()
    sog = next((b for b in piano.get(1, []) if b[0] == "tabella" and b[1].startswith("Soggetti")), None)
    committente = next((r[1] for r in (sog[3] if sog else []) if "committente" in r[0].lower()), DAV)
    nome_csp = (f"{coordinatore.get('nome', '')} {coordinatore.get('cognome', '')}".strip() if coordinatore else "") or DAV
    _tabella(doc, ["", ""], [["Cantiere", p["nome"]], ["Indirizzo", p.get("indirizzo_cantiere") or DAV],
                             ["Committente", committente], ["Coordinatore per la progettazione", nome_csp],
                             ["Revisione", rev]], dim=10, larghezze=[5.0, 12.0], intestazione=False)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ── Punti da verificare e indice ──
    punti = punti_da_verificare(piano)
    doc.add_heading("Punti da verificare", level=1)
    if punti:
        p0 = doc.add_paragraph()
        _run_dav(p0, f"Il documento contiene {len(punti)} parti con dati DA VERIFICARE, evidenziati in giallo nel testo. "
                     "Vanno completati prima dell'emissione definitiva.")
        _tabella(doc, ["Cap.", "Capitolo", "Dove"], [[str(n), t, d] for n, t, d in punti], larghezze=[1.2, 5.0, 10.8])
    else:
        doc.add_paragraph("Nessun dato da verificare.")
    doc.add_heading("Indice", level=1)
    for n, titolo in CAPITOLI:
        doc.add_paragraph(f"{n}. {titolo}")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ── Capitoli ──
    for n, titolo in CAPITOLI:
        doc.add_heading(f"{n}. {titolo}", level=1)
        blocchi = piano.get(n, [])
        if not blocchi:
            _run_dav(doc.add_paragraph(), "Contenuto non disponibile: DA VERIFICARE.")
        for b in blocchi:
            if b[0] == "testo":
                _sottotitolo(doc, b[1])
                for riga in (b[2] or "").split("\n"):
                    if riga.strip():
                        _run_dav(doc.add_paragraph(), riga.strip())
            elif b[0] == "tabella":
                _sottotitolo(doc, b[1])
                _tabella(doc, b[2], b[3])
            elif b[0] == "gantt":
                _sottotitolo(doc, "Diagramma di Gantt (settimane)")
                _gantt(doc, b[1], b[2])
            elif b[0] == "costi":
                _sottotitolo(doc, "Stima analitica dei costi della sicurezza (Allegato XV, punto 4.1.1)")
                _costi(doc, b[1], b[2])
            elif b[0] == "immagine":
                _sottotitolo(doc, b[1])
                import io
                doc.add_paragraph().add_run().add_picture(io.BytesIO(b[2]), width=Cm(LARGHEZZA_UTILE_CM))
            elif b[0] == "nota":
                _run_dav(doc.add_paragraph(), b[1])

    # ── Firma ──
    doc.add_paragraph()
    f = doc.add_paragraph(); f.paragraph_format.keep_with_next = True
    _run_dav(f, f"Luogo e data: ____________________, {data_rev}")
    f = doc.add_paragraph(); f.alignment = WD_ALIGN_PARAGRAPH.RIGHT; f.paragraph_format.keep_with_next = True
    r = f.add_run("Il Coordinatore per la sicurezza in fase di progettazione"); r.bold = True
    if coordinatore:
        f = doc.add_paragraph(); f.alignment = WD_ALIGN_PARAGRAPH.RIGHT; f.paragraph_format.keep_with_next = True
        titolo_cs = " ".join(x for x in (coordinatore.get("titolo_studio"), coordinatore.get("nome"), coordinatore.get("cognome")) if x)
        f.add_run(titolo_cs)
    f = doc.add_paragraph(); f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if firma_path and os.path.exists(firma_path):
        f.add_run().add_picture(firma_path, height=Cm(2.0))
    else:
        doc.add_paragraph()
        f.add_run("_________________________________")

    import io
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ══════════════════════════════════════════════════════════════════════════════
# Export con salvataggio nello Storico
# ══════════════════════════════════════════════════════════════════════════════

def esporta(progetto_id: int, user: dict, revisione: int, data_rev: str, coordinatore: dict, studio: dict) -> dict:
    dati = costruisci_piano(progetto_id, user["username"], coordinatore, studio)
    contenuto = scrivi(dati, studio, coordinatore, (coordinatore or {}).get("firma_path"), studio.get("logo_path"),
                       revisione, data_rev)
    base = re.sub(r"[^A-Za-z0-9]+", "_", dati["progetto"]["nome"]).strip("_")[:40] or "PSC"
    path = os.path.join(cartella_documenti(), f"PSC_{base}_Rev{revisione}_{uuid.uuid4().hex[:6]}.docx")
    with open(path, "wb") as fh:
        fh.write(contenuto)
    conn = get_conn()
    try:
        esistente = conn.execute(
            "SELECT id, file_path FROM documenti WHERE progetto_id=? AND revisione=? AND tipo='psc' AND username=?",
            (progetto_id, revisione, user["username"])).fetchone()
        meta = json.dumps({"progetto_id": progetto_id, "revisione": revisione, "data": data_rev}, ensure_ascii=False)
        nome = f"{dati['progetto']['nome']} — PSC Rev. {revisione}"
        if esistente:
            # Stessa revisione già nello Storico: il file viene sostituito con quello nuovo
            if esistente["file_path"] and os.path.exists(esistente["file_path"]):
                os.remove(esistente["file_path"])
            conn.execute("UPDATE documenti SET file_path=?, contenuto=?, nome_cantiere=?, created_at=CURRENT_TIMESTAMP WHERE id=?",
                         (path, meta, nome, esistente["id"]))
            doc_id, azione = esistente["id"], "sostituita"
        else:
            cur = conn.execute(
                """INSERT INTO documenti (tipo, nome_cantiere, contenuto, stato, username, file_path, progetto_id, revisione)
                   VALUES ('psc', ?, ?, 'completato', ?, ?, ?, ?)""",
                (nome, meta, user["username"], path, progetto_id, revisione))
            doc_id, azione = cur.lastrowid, "nuova"
        conn.commit()
    finally:
        conn.close()
    return {"doc_id": doc_id, "revisione": revisione, "storico": azione,
            "punti_da_verificare": len(punti_da_verificare(dati["piano"]))}
